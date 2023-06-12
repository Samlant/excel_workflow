import os
import time
import subprocess
import shutil
import string
from tkinter import *
from datetime import datetime
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
import fillpdf
from fillpdf import fillpdfs


# # Assuming CURRENT WORKING DIR is USER\APPDATA\LOCAL\SAM_PROGRAM
# HOME_DIR = os.path.expanduser( '~' )
# PATH_TO_WATCH = os.path.join(HOME_DIR, 'Novamar FL Office Drive') # Root of FL Office Shared Drive
# QUOTE_FOLDER = os.path.join(PATH_TO_WATCH, 'QUOTES New')
# TRACKER_PATH = os.path.join(PATH_TO_WATCH, 'Trackers')


PATH_TO_WATCH = os.getcwd()
QUOTES_FOLDER = os.path.join(PATH_TO_WATCH, "QUOTES New")
TRACKER_PATH = os.path.join(
    PATH_TO_WATCH,
    "Trackers",
    "1MASTER 2023 QUOTE TRACKER.xlsx",
)


class DirWatch:
    def __init__(self):
        self._begin_watch()

    def _begin_watch(self) -> None:
        before = dict([(f, None) for f in os.listdir(PATH_TO_WATCH)])
        while 1:
            time.sleep(2)
            after = dict([(f, None) for f in os.listdir(PATH_TO_WATCH)])
            added = [f for f in after if not f in before]
            if added:
                new_file = added[0]
                dialog = DialogNewFile(new_file)
                dialog.root.mainloop()
                before = dict([(f, None) for f in os.listdir(PATH_TO_WATCH)])


class DialogNewFile:
    def __init__(self, file_name):
        self.excel_entry = {
            "vessel_year": "",
            "vessel": "",
            "markets": {},
            "status": "ALLOCATE AND SUBMIT TO MRKTS",
            "referral": "",
        }
        self.file_name = file_name
        self._initialize()

    def _initialize(self):
        self.root = Tk()
        self.root.geometry("300x400")
        self.root.title("Next Steps")
        self.root.text_frame = Frame(self.root, bg="#CFEBDF")
        self.root.text_frame.pack(fill=BOTH, expand=True)
        self.root.btn_frame = Frame(self.root, bg="#CFEBDF")
        self.root.btn_frame.pack(fill=BOTH, expand=True, ipady=2)
        self._save_client_name()
        if os.path.splitext(self.file_name)[1] == ".pdf":
            self.get_PDF_values()
        self._create_widgets()

    def _save_client_name(self) -> None:
        client_name = os.path.splitext(self.file_name)[0].split(" ")
        self.excel_entry["fname"] = string.capwords(client_name[1])
        self.excel_entry["lname"] = client_name[0].upper()

    def get_PDF_values(self):
        keys_dict = {
            "fname": "fname",
            "lname": "lname",
            "year": "vessel_year",
            "vessel": "vessel_make_model",
            "referral": "referral",
        }
        pdf_dict = fillpdfs.get_form_fields(self.file_name)
        pdf_dict = {key: pdf_dict[key] for key in pdf_dict.keys() & keys_dict.values()}

        self.excel_entry = {}
        fname = pdf_dict.get(keys_dict["fname"])
        self.excel_entry["fname"] = string.capwords(fname)
        lname = pdf_dict.get(keys_dict["lname"])
        self.excel_entry["lname"] = lname.upper()
        vessel = pdf_dict.get(keys_dict["vessel"])
        self.excel_entry["vessel"] = string.capwords(vessel)
        self.excel_entry["vessel_year"] = pdf_dict.get(keys_dict["year"])
        referral = pdf_dict.get(keys_dict["referral"])
        self.excel_entry["referral"] = referral.upper()
        # if any(chr.isdigit() for chr in self.excel_entry["vessel"]):
        #     self.excel_entry["length"] = pdf_dict.get(keys_dict["length"])

    def _create_widgets(self):
        client_name = " ".join([self.excel_entry["fname"], self.excel_entry["lname"]])
        vessel = self.excel_entry["vessel"]
        year = self.excel_entry["vessel_year"]
        referral = self.excel_entry["referral"]
        self.root.text_frame.grid_columnconfigure(0, weight=1)
        self.root.btn_frame.grid_columnconfigure(0, weight=1)

        Label(self.root.text_frame, text="Client name: ", bg="#CFEBDF").grid(
            column=0, row=0, pady=(3, 0)
        )
        name_entry = Entry(
            self.root.text_frame, width=30, justify="center", bg="#5F634F", fg="#FFCAB1"
        )
        name_entry.insert(0, client_name)
        name_entry.grid(column=0, row=1, pady=(0, 8))

        Label(self.root.text_frame, text="Vessel: ", bg="#CFEBDF").grid(column=0, row=2)
        vessel_entry = Entry(
            self.root.text_frame, width=30, justify="center", bg="#5F634F", fg="#FFCAB1"
        )
        vessel_entry.insert(0, vessel)
        vessel_entry.grid(column=0, row=3, pady=(0, 8))

        Label(self.root.text_frame, text="Vessel year: ", bg="#CFEBDF").grid(
            column=0, row=4
        )
        year_entry = Entry(
            self.root.text_frame, width=10, justify="center", bg="#5F634F", fg="#FFCAB1"
        )
        year_entry.insert(0, year)
        year_entry.grid(column=0, row=5, pady=(0, 8))

        Label(self.root.text_frame, text="Referral: ", bg="#CFEBDF").grid(
            column=0, row=6
        )
        referral_entry = Entry(
            self.root.text_frame,
            width=30,
            justify="center",
            bg="#5F634F",
            fg="#FFCAB1",
        )
        referral_entry.insert(0, referral)
        referral_entry.grid(column=0, row=7, pady=(0, 7))

        submit_btn = Button(
            self.root.btn_frame,
            text="Submit to Markets",
            width=36,
            height=3,
            command=lambda: self.choice("submit"),
            default=ACTIVE,
            bg="#1D3461",
            fg="#CFEBDF",
        )
        submit_btn.grid(row=0, column=0, padx=5, pady=(0, 0))

        allocate_btn = Button(
            self.root.btn_frame,
            text="Allocate Markets",
            width=36,
            height=3,
            command=lambda: self.choice("allocate"),
            default=ACTIVE,
            bg="#1D3461",
            fg="#CFEBDF",
        )
        allocate_btn.grid(row=1, column=0, padx=5, pady=(3, 3))

        create_folder_only_btn = Button(
            self.root.btn_frame,
            text="Only create folder",
            width=36,
            height=3,
            command=lambda: self.choice("only create folder"),
            default=ACTIVE,
            bg="#1D3461",
            fg="#CFEBDF",
        )
        create_folder_only_btn.grid(row=2, column=0, padx=5, pady=(0, 5))

    def choice(self, option: str) -> None:
        if option == "only create folder":
            self._create_folder()
            self.root.destroy()
            self._create_excel_entry()

        elif option == "allocate":
            self._create_folder()
            self.root.destroy()
            self.allocate_markets()

        else:
            self._create_folder()
            self.root.destroy()
            self.run_quickdraw_app()

    def _create_folder(self):
        self.dir_name = os.path.splitext(self.file_name)[0]
        # dir_name = dir_name.split() #NOT needed FOR NOW as we will title files with client names ... for now
        path = os.path.join(QUOTES_FOLDER, self.dir_name)
        os.makedirs(path)
        self._move_quoteform_to_folder(path)

    def _move_quoteform_to_folder(self, path: str):
        shutil.move(self.file_name, path)

    def _create_excel_entry(self):
        excel = ExcelWorker(self.excel_entry)
        row_data = excel.create_entry_list()
        excel.create_row(row_data=row_data)
        excel.save_workbook()

    def allocate_markets(self) -> dict:
        dialog_allocate = DialogAllocateMarkets()
        self.excel_entry["markets"] = dialog_allocate._return_markets()
        self._create_excel_entry()

    def run_quickdraw_app(self):
        path = os.path.join(HOME_DIR, "AppData", "Sam_Programs", "QuickDraw.exe")
        subprocess.run([path])


class DialogAllocateMarkets:
    def __init__(self):
        self._initialize()

    def _initialize(self):
        self.root = Tk()
        self.root.geometry("260x560")
        self.root.title("Allocate Markets")
        self.root.frame = Frame(self.root, bg="#CFEBDF")
        self.root.frame.pack(fill=BOTH, expand=False)
        self._create_widgets()

    def _create_widgets(self):
        Label(
            self.root.frame,
            text="ALLOCATE MARKETS",
            justify="center",
            bg="#CFEBDF",
            fg="#5F634F",
        ).pack(fill=X, ipady=6)
        self.ch_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="Chubb",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.ch_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )
        self.mk_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="Markel",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.mk_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )
        self.ai_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="American Integrity",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.ai_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )
        self.am_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="American Modern",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.am_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )
        self.pg_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="Progressive",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.pg_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )
        self.sw_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="Seawave",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.sw_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )
        self.km_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="Kemah Marine",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.km_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )
        self.cp_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="Concept Special Risks",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.cp_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )
        self.nh_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="New Hampshire",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.nh_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )
        self.In_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="Intact",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.In_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )
        self.tv_checkbtn = Checkbutton(
            self.root.frame,
            relief="raised",
            text="Travelers",
            justify=CENTER,
            anchor=W,
            fg="#FFCAB1",
            bg="#5F634F",
            selectcolor="#000000",
        )
        self.tv_checkbtn.pack(
            fill=X, expand=False, ipady=6, ipadx=10, pady=3, padx=10, anchor=NW
        )

        allocate_btn = Button(
            master=self.root.frame,
            text="ALLOCATE",
            width=30,
            height=10,
            bg="#1D3461",
            fg="#CFEBDF",
        )
        allocate_btn.pack(
            fill=X,
            expand=False,
            pady=5,
            padx=10,
            ipady=6,
            ipadx=10,
            command=lambda: self._return_markets(),
        )

    def _return_markets(self):
        dict_of_markets = {
            "ch": self.ch_checkbtn.get(),
            "mk": self.mk_checkbtn.get(),
            "ai": self.ai_checkbtn.get(),
            "am": self.am_checkbtn.get(),
            "pg": self.pg_checkbtn.get(),
            "sw": self.sw_checkbtn.get(),
            "km": self.km_checkbtn.get(),
            "cp": self.cp_checkbtn.get(),
            "nh": self.nh_checkbtn.get(),
            "In": self.In_checkbtn.get(),
            "tv": self.tv_checkbtn.get(),
        }
        return dict_of_markets


class ExcelWorker:
    def __init__(
        self,
        excel_entry: dict,
    ):
        fname = excel_entry["fname"].capitalize()
        lname = excel_entry["lname"].upper()
        self.name = " ".join([lname, fname])
        self.date = str(datetime.today()).split()[0]
        self.vessel_year = excel_entry["vessel_year"]
        self.vessel = excel_entry["vessel"]
        self.markets = excel_entry["markets"]
        self.status = excel_entry["status"]
        self.referral = excel_entry["referral"]
        self.wb = load_workbook(TRACKER_PATH)
        month = self.get_current_month()
        self.ws = self.wb[month]
        self.markets = self._assign_markets()

    def get_current_month(self):
        months_of_the_year = {
            1: "January",
            2: "February",
            3: "March",
            4: "April",
            5: "May",
            6: "June",
            7: "July",
            8: "August",
            9: "September",
            10: "October",
            11: "November",
            12: "December",
        }
        month = datetime.now().month
        return months_of_the_year.get(month).upper()

    def _assign_markets(self) -> list:
        list_of_markets = []
        for key, value in self.markets:
            if value == 1:
                list_of_markets.append(value)
            else:
                list_of_markets.append("")
        return list_of_markets

    def create_entry_list(self) -> list:
        list_of_client_data = (
            [
                "",
                "",
                "",
                self.name,
                self.date,
                "",
                self.vessel_year,
                self.vessel,
            ]
            + self.markets
            + [
                "",
                "",
                "",
                self.status,
                self.referral,
            ]
        )
        return list_of_client_data

    def create_row(self, row_data: list):
        if self.ws.append(row_data):
            if self._save_workbook():
                return True
            else:
                return False
        else:
            return False

    def save_workbook(self):
        self.wb.save(self.wb_path)


app = DirWatch()
